#!/usr/bin/env python3
"""Convert a DOCX document to a PPTX presentation.

Usage:
    python convert_docx_to_pptx.py input.docx [output.pptx]

If output path is omitted the PPTX file is written next to the DOCX file
with the same base name.
"""

import sys
import os
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.33)
SLIDE_HEIGHT = Inches(7.5)

TITLE_FONT_SIZE = Pt(36)
HEADING_FONT_SIZE = Pt(28)
BODY_FONT_SIZE = Pt(18)
MAX_CHARS_PER_SLIDE = 600


def _is_heading(paragraph):
    """Return a tuple (bool, int): True/level if heading, otherwise (False, 0)."""
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("Heading"):
        try:
            level = int(style_name.split()[-1])
        except ValueError:
            level = 1
        return True, level
    return False, 0


def _paragraph_text(paragraph):
    return paragraph.text.strip()


def _add_title_slide(prs, title_text, subtitle_text=""):
    layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title_text
    if subtitle_text and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle_text
    return slide


def _add_content_slide(prs, title_text, body_lines):
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title_text

    tf = slide.placeholders[1].text_frame
    tf.word_wrap = True
    tf.clear()

    first = True
    for line in body_lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = line
        if p.runs:
            p.runs[0].font.size = BODY_FONT_SIZE

    return slide


def convert(docx_path, pptx_path):
    doc = Document(docx_path)
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Gather paragraphs with metadata
    paragraphs = []
    for para in doc.paragraphs:
        text = _paragraph_text(para)
        if not text:
            continue
        is_h, level = _is_heading(para)
        paragraphs.append({"text": text, "is_heading": is_h, "level": level})

    if not paragraphs:
        print("No content found in the document.")
        return

    # Use the first Heading 1 (or first paragraph) as the title slide
    title_text = None
    start_idx = 0
    for i, p in enumerate(paragraphs):
        if p["is_heading"] and p["level"] == 1:
            title_text = p["text"]
            start_idx = i + 1
            break
    if title_text is None:
        title_text = paragraphs[0]["text"]
        start_idx = 1

    # Determine subtitle (use next heading or first body paragraph, but don't skip body paragraphs)
    subtitle_text = ""
    if start_idx < len(paragraphs):
        subtitle_text = paragraphs[start_idx]["text"]
        # Only advance past the subtitle if it is a heading; body text must remain in content slides
        if paragraphs[start_idx]["is_heading"]:
            start_idx += 1

    _add_title_slide(prs, title_text, subtitle_text)

    # Build content slides grouped by headings
    current_heading = title_text
    body_buffer = []
    char_count = 0

    def flush_slide():
        nonlocal body_buffer, char_count
        if body_buffer:
            _add_content_slide(prs, current_heading, body_buffer)
        body_buffer = []
        char_count = 0

    for para in paragraphs[start_idx:]:
        text = para["text"]
        if para["is_heading"]:
            flush_slide()
            current_heading = text
        else:
            if char_count + len(text) > MAX_CHARS_PER_SLIDE and body_buffer:
                flush_slide()
            body_buffer.append(text)
            char_count += len(text)

    flush_slide()

    prs.save(pptx_path)
    print(f"Saved presentation to: {pptx_path}")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    docx_path = sys.argv[1]
    if not os.path.isfile(docx_path):
        print(f"Error: file not found: {docx_path}")
        sys.exit(1)

    if len(sys.argv) >= 3:
        pptx_path = sys.argv[2]
    else:
        base = os.path.splitext(docx_path)[0]
        pptx_path = base + ".pptx"

    convert(docx_path, pptx_path)


if __name__ == "__main__":
    main()
